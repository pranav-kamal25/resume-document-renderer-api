from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .models import ResumeContent


# Resume PDF format constants.
FONT = "Carlito"

PAGE_W_PT = 612.0
PAGE_H_PT = 792.0

GOLDEN_MAX_TEXT_Y_PT = 757.864013671875

# Allow modest content-dependent variance.
GOLDEN_TOLERANCE_PT = 28.0


def _set_run(
    run,
    bold=None,
    size=10.5,
):
    run.font.name = FONT
    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FONT,
    )
    run.font.size = Pt(size)
    run.bold = bold


def _safe_text(text: str) -> str:
    # Keep candidate wording intact except for known rendering hazards.
    text = (
        text.replace("\u00a0", " ")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )

    if (
        "\ufffd" in text
        or "\ufffe" in text
        or "\uffff" in text
    ):
        raise ValueError(
            "Input contains malformed/replacement Unicode characters"
        )

    return text


def _para(
    doc,
    text="",
    before=0,
    after=1,
    align=None,
    bold=False,
    size=10.5,
):
    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(
        before
    )
    paragraph.paragraph_format.space_after = Pt(
        after
    )
    paragraph.paragraph_format.line_spacing = 1.0

    if align is not None:
        paragraph.alignment = align

    run = paragraph.add_run(
        _safe_text(text)
    )
    _set_run(
        run,
        bold,
        size,
    )

    return paragraph


def _labeled(
    doc,
    label,
    text,
    before=0,
    after=1,
):
    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(
        before
    )
    paragraph.paragraph_format.space_after = Pt(
        after
    )
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(
        _safe_text(label)
    )
    _set_run(
        run,
        True,
    )

    run = paragraph.add_run(
        _safe_text(text)
    )
    _set_run(
        run,
        False,
    )

    return paragraph


def _bullet(
    doc,
    text,
    after=1,
    before=0,
):
    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(
        before
    )
    paragraph.paragraph_format.space_after = Pt(
        after
    )
    paragraph.paragraph_format.line_spacing = 1.0

    # FINAL LOCKED hanging indent:
    # all continuation lines align with the first text character.
    text_x = Inches(0.2375)

    paragraph.paragraph_format.left_indent = text_x
    paragraph.paragraph_format.first_line_indent = Inches(
        -0.115
    )

    p_pr = paragraph._p.get_or_add_pPr()

    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")

    tab.set(
        qn("w:val"),
        "left",
    )
    tab.set(
        qn("w:pos"),
        str(
            int(
                0.2375 * 1440
            )
        ),
    )

    tabs.append(tab)
    p_pr.append(tabs)

    run = paragraph.add_run(
        "•\t"
    )
    _set_run(run)

    run = paragraph.add_run(
        _safe_text(text)
    )
    _set_run(run)

    return paragraph


def _section(
    doc,
    title,
):
    table = doc.add_table(
        rows=1,
        cols=1,
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )
    table.autofit = False

    table.columns[0].width = Inches(
        7.26
    )

    cell = table.cell(
        0,
        0,
    )

    cell.width = Inches(
        7.26
    )

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    tc_pr = cell._tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in(
        "w:tcMar"
    )

    if tc_mar is None:
        tc_mar = OxmlElement(
            "w:tcMar"
        )
        tc_pr.append(
            tc_mar
        )

    for side in [
        "top",
        "left",
        "bottom",
        "right",
    ]:
        element = tc_mar.find(
            qn(
                "w:" + side
            )
        )

        if element is None:
            element = OxmlElement(
                "w:" + side
            )
            tc_mar.append(
                element
            )

        element.set(
            qn("w:w"),
            "0",
        )
        element.set(
            qn("w:type"),
            "dxa",
        )

    borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if borders is None:
        borders = OxmlElement(
            "w:tcBorders"
        )
        tc_pr.append(
            borders
        )

    for edge in [
        "top",
        "bottom",
    ]:
        element = borders.find(
            qn(
                "w:" + edge
            )
        )

        if element is None:
            element = OxmlElement(
                "w:" + edge
            )
            borders.append(
                element
            )

        element.set(
            qn("w:val"),
            "single",
        )
        element.set(
            qn("w:sz"),
            "6",
        )
        element.set(
            qn("w:space"),
            "0",
        )
        element.set(
            qn("w:color"),
            "666666",
        )

    for edge in [
        "left",
        "right",
        "insideH",
        "insideV",
    ]:
        element = borders.find(
            qn(
                "w:" + edge
            )
        )

        if element is None:
            element = OxmlElement(
                "w:" + edge
            )
            borders.append(
                element
            )

        element.set(
            qn("w:val"),
            "nil",
        )

    paragraph = cell.paragraphs[0]

    paragraph.paragraph_format.space_before = Pt(
        0
    )
    paragraph.paragraph_format.space_after = Pt(
        0
    )
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(
        _safe_text(title)
    )

    _set_run(
        run,
        True,
        13.5,
    )

    return table


def build_docx(
    content: ResumeContent,
    out_path: Path,
) -> None:
    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(
        0.42
    )
    section.bottom_margin = Inches(
        0.30
    )
    section.left_margin = Inches(
        0.62
    )
    section.right_margin = Inches(
        0.62
    )
    section.header_distance = Inches(
        0.2
    )
    section.footer_distance = Inches(
        0.2
    )

    style = doc.styles["Normal"]

    style.font.name = FONT
    style.font.size = Pt(
        10.5
    )

    style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FONT,
    )

    paragraph_format = style.paragraph_format

    paragraph_format.space_before = Pt(
        0
    )
    paragraph_format.space_after = Pt(
        1
    )
    paragraph_format.line_spacing = 1.0

    header = content.header

    _para(
        doc,
        header.name,
        0,
        0.5,
        WD_ALIGN_PARAGRAPH.CENTER,
        True,
        20,
    )

    contact = (
        f"{header.location} | "
        f"{header.email} | "
        f"{header.phone} | "
        f"{header.linkedin}"
    )

    _para(
        doc,
        contact,
        0,
        3,
        WD_ALIGN_PARAGRAPH.CENTER,
        False,
        10.5,
    )

    _section(
        doc,
        "EDUCATION",
    )

    _para(
        doc,
        content.education.school_line,
        3,
        0.5,
        bold=True,
    )

    _labeled(
        doc,
        "GPA: ",
        content.education.gpa_text,
        0,
        1,
    )

    if content.education.coursework:
        _labeled(
            doc,
            "Relevant Coursework: ",
            ", ".join(
                map(
                    _safe_text,
                    content.education.coursework,
                )
            ),
            0,
            2,
        )

    if content.experiences:
        _section(
            doc,
            "EXPERIENCE",
        )

        for index, experience in enumerate(
            content.experiences
        ):
            _para(
                doc,
                experience.heading,
                3 if index == 0 else 2,
                0.5,
                bold=True,
            )

            for bullet_index, bullet in enumerate(
                experience.bullets
            ):
                _bullet(
                    doc,
                    bullet,
                    after=(
                        2
                        if bullet_index
                        == len(
                            experience.bullets
                        )
                        - 1
                        else 1
                    ),
                )

    if content.projects:
        _section(
            doc,
            "PROJECTS",
        )

        for index, project in enumerate(
            content.projects
        ):
            _para(
                doc,
                project.heading,
                3 if index == 0 else 2,
                0.5,
                bold=True,
            )

            for bullet_index, bullet in enumerate(
                project.bullets
            ):
                _bullet(
                    doc,
                    bullet,
                    after=(
                        2
                        if bullet_index
                        == len(
                            project.bullets
                        )
                        - 1
                        else 1
                    ),
                )

    if content.certifications:
        _section(
            doc,
            "CERTIFICATIONS & ONLINE COURSEWORK",
        )

        for index, certification in enumerate(
            content.certifications
        ):
            _bullet(
                doc,
                certification,
                after=(
                    2
                    if index
                    == len(
                        content.certifications
                    )
                    - 1
                    else 1
                ),
                before=(
                    3
                    if index == 0
                    else 0
                ),
            )

    if content.skills:
        _section(
            doc,
            "SKILLS",
        )

        for index, skill_group in enumerate(
            content.skills
        ):
            _labeled(
                doc,
                f"{skill_group.label}: ",
                ", ".join(
                    map(
                        _safe_text,
                        skill_group.items,
                    )
                ),
                3 if index == 0 else 0,
                1,
            )

    out_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    doc.save(
        str(
            out_path
        )
    )


def convert_to_pdf(
    docx_path: Path,
    out_dir: Path,
) -> Path:
    out_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    profile = tempfile.mkdtemp(
        prefix="lo-profile-"
    )

    try:
        command = [
            "libreoffice",
            "--headless",
            (
                "-env:UserInstallation="
                f"file://{profile}"
            ),
            "--convert-to",
            "pdf",
            "--outdir",
            str(
                out_dir
            ),
            str(
                docx_path
            ),
        ]

        result = subprocess.run(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=60,
        )

        pdf_path = (
            out_dir
            / (
                docx_path.stem
                + ".pdf"
            )
        )

        if (
            result.returncode != 0
            or not pdf_path.exists()
        ):
            raise RuntimeError(
                "LibreOffice PDF conversion failed: "
                + (
                    result.stderr
                    or result.stdout
                )
            )

        return pdf_path

    finally:
        shutil.rmtree(
            profile,
            ignore_errors=True,
        )


def inspect_pdf(
    pdf_path: Path,
) -> dict:
    warnings = []

    try:
        document = fitz.open(
            pdf_path
        )
    except Exception as exc:
        return {
            "status": "FAIL",
            "page_count": 0,
            "max_text_y_pt": None,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": None,
            "fit_direction": "FIX_TECHNICAL",
            "warnings": [
                (
                    "Unable to inspect generated PDF: "
                    f"{exc}"
                )
            ],
        }

    page_count = len(
        document
    )

    if page_count == 0:
        return {
            "status": "FAIL",
            "page_count": 0,
            "max_text_y_pt": None,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": None,
            "fit_direction": "FIX_TECHNICAL",
            "warnings": [
                "Generated PDF contains no pages"
            ],
        }

    ys = []

    malformed = False

    for page in document:
        text = page.get_text(
            "text"
        )

        if (
            "\ufffd" in text
            or "\ufffe" in text
            or "\uffff" in text
        ):
            malformed = True
            warnings.append(
                "Malformed/replacement character detected"
            )

        for block in page.get_text(
            "blocks"
        ):
            (
                _x0,
                _y0,
                _x1,
                y1,
                text_block,
                *_rest,
            ) = block

            if text_block.strip():
                ys.append(
                    float(
                        y1
                    )
                )

    max_y = (
        max(ys)
        if ys
        else None
    )

    delta = (
        None
        if max_y is None
        else (
            max_y
            - GOLDEN_MAX_TEXT_Y_PT
        )
    )

    if malformed:
        return {
            "status": "FAIL",
            "page_count": page_count,
            "max_text_y_pt": max_y,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": delta,
            "fit_direction": "FIX_TECHNICAL",
            "warnings": warnings,
        }

    if max_y is None:
        warnings.append(
            "No text blocks detected"
        )

        return {
            "status": "FAIL",
            "page_count": page_count,
            "max_text_y_pt": None,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": None,
            "fit_direction": "FIX_TECHNICAL",
            "warnings": warnings,
        }

    # More than one page is a CONTENT-FIT issue,
    # not a renderer/system failure.
    if page_count > 1:
        warnings.append(
            (
                "OVERFLOW: expected exactly 1 page, "
                f"got {page_count}"
            )
        )

        return {
            "status": "OVERFLOW",
            "page_count": page_count,
            "max_text_y_pt": max_y,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": delta,
            "fit_direction": "REMOVE_CONTENT",
            "warnings": warnings,
        }

    if (
        max_y
        < (
            GOLDEN_MAX_TEXT_Y_PT
            - GOLDEN_TOLERANCE_PT
        )
    ):
        warnings.append(
            (
                "UNDERFILL: content ends materially "
                "above the locked golden reference"
            )
        )

        return {
            "status": "UNDERFILL",
            "page_count": page_count,
            "max_text_y_pt": max_y,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": delta,
            "fit_direction": "ADD_CONTENT",
            "warnings": warnings,
        }

    if (
        max_y
        > (
            GOLDEN_MAX_TEXT_Y_PT
            + GOLDEN_TOLERANCE_PT
        )
    ):
        warnings.append(
            (
                "OVERDENSE: content extends materially "
                "below the locked golden reference"
            )
        )

        return {
            "status": "OVERDENSE",
            "page_count": page_count,
            "max_text_y_pt": max_y,
            "golden_max_text_y_pt": (
                GOLDEN_MAX_TEXT_Y_PT
            ),
            "bottom_delta_from_golden_pt": delta,
            "fit_direction": "REMOVE_CONTENT",
            "warnings": warnings,
        }

    return {
        "status": "PASS",
        "page_count": page_count,
        "max_text_y_pt": max_y,
        "golden_max_text_y_pt": (
            GOLDEN_MAX_TEXT_Y_PT
        ),
        "bottom_delta_from_golden_pt": delta,
        "fit_direction": "NONE",
        "warnings": warnings,
    }
